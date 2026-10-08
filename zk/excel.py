"""Son durumdan Excel raporu (Özet / Kampanyalar / Geçmiş)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def _g(d, *yol):
    for k in yol:
        if d is None:
            return None
        d = d.get(k)
    return d


OZET = [
    ("Marka", lambda v: v["marka"]),
    ("Aktif Kampanya (web)", lambda v: len(v["kampanyalar"]) if v["kampanyalar"] is not None else "-"),
    ("App Store Puan", lambda v: _g(v, "appstore", "puan")),
    ("App Store Oy", lambda v: _g(v, "appstore", "oy")),
    ("Google Play Puan", lambda v: _g(v, "play", "puan")),
    ("Google Play Oy", lambda v: _g(v, "play", "oy")),
    ("Google Play İndirme", lambda v: _g(v, "play", "indirme")),
    ("Şikayet (son 1 yıl)", lambda v: _g(v, "sikayetvar", "sikayet_1yil")),
    ("Çözülen (son 1 yıl)", lambda v: _g(v, "sikayetvar", "cozulen_1yil")),
    ("Çözüm Başarısı % (son 1 yıl)", lambda v: _g(v, "sikayetvar", "cozum_orani_1yil")),
    ("Şikayet (son 1 ay)", lambda v: _g(v, "sikayetvar", "sikayet_1ay")),
    ("Şikayet (tüm zamanlar)", lambda v: _g(v, "sikayetvar", "sikayet_tum")),
]


def _baslik(ws, basliklar):
    ws.append(basliklar)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4B2E83")
    ws.freeze_panes = "A2"


def _genislik(ws):
    for i, kol in enumerate(ws.columns, 1):
        uz = max(len(str(c.value or "")) for c in kol)
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, uz + 2), 90)


def rapor_olustur(d, hedef):
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    _baslik(ws, [b for b, _ in OZET] + ["Yeni Kampanya", "Kalkan Kampanya"])
    for v in d["markalar"]:
        f = d["fark"][v["marka"]]
        ws.append([fn(v) for _, fn in OZET] + [len(f["yeni"]), len(f["kalkan"])])
    ws.append([])
    ws.append([f"Güncelleme: {d['guncelleme']} UTC" + (f" · Karşılaştırma: {d['onceki_tarih']}" if d["onceki_tarih"] else "")])
    _genislik(ws)

    ws = wb.create_sheet("Kampanyalar")
    _baslik(ws, ["Marka", "Kampanya", "Durum", "Link"])
    yesil, gri = PatternFill("solid", fgColor="D9F2D9"), PatternFill("solid", fgColor="EEEEEE")
    for v in d["markalar"]:
        f = d["fark"][v["marka"]]
        if v["kampanyalar"] is None:
            ws.append([v["marka"], "Web sitesinde kampanya listesi yok / çekilemedi", "", ""])
            continue
        for k in v["kampanyalar"]:
            yeni = k["link"] in f["yeni"]
            ws.append([v["marka"], k["baslik"], "YENİ" if yeni else "", k["link"]])
            if yeni:
                for c in ws[ws.max_row]:
                    c.fill = yesil
        for k in f["kalkan"]:
            ws.append([v["marka"], k["baslik"], "KALKTI", k["link"]])
            for c in ws[ws.max_row]:
                c.fill = gri
    ws.auto_filter.ref = ws.dimensions
    _genislik(ws)

    ws = wb.create_sheet("Geçmiş")
    _baslik(ws, ["Tarih", "Marka", "Aktif Kampanya", "App Store", "Google Play", "Şikayet (1 yıl)", "Çözüm % (1 yıl)"])
    for gun in d["gecmis"]:
        for marka, x in gun["markalar"].items():
            ws.append([gun["tarih"], marka, x["kampanya"], x["appstore"], x["play"], x["sikayet_1yil"], x["cozum_orani_1yil"]])
    ws.auto_filter.ref = ws.dimensions
    _genislik(ws)

    wb.save(hedef)
